#!/usr/bin/env python3
"""生成 geminus 测试夹具 DOCX（docs/superpowers/plans 任务 T2）。

输出：
  test/fixtures/simple.docx —— 标题/正文/混排 run/列表/空段/居中
  test/fixtures/rich.docx   —— simple 全部 + 块级图片 + 行内图片 + 3x3 表格 + OMML 公式（独立段 + 行内）
"""
import io
import os
import struct
import zlib

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.shared import Inches

OUT_DIR = os.path.join(os.path.dirname(__file__), "..", "test", "fixtures")

M_NS = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'


def make_png(width: int, height: int, rgb: tuple) -> bytes:
    """纯 Python 生成最小 PNG（RGB，无压缩依赖）。"""
    def chunk(typ: bytes, data: bytes) -> bytes:
        c = struct.pack(">I", len(data)) + typ + data
        c += struct.pack(">I", zlib.crc32(typ + data) & 0xFFFFFFFF)
        return c

    ihdr = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    row = b"\x00" + bytes(rgb) * width
    raw = row * height
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


def omml_e_mc2(display: bool) -> str:
    inner = (
        "<m:sSup>"
        "<m:e><m:r><m:t>E=mc</m:t></m:r></m:e>"
        "<m:sup><m:r><m:t>2</m:t></m:r></m:sup>"
        "</m:sSup>"
    )
    math = f"<m:oMath {M_NS}>{inner}</m:oMath>"
    if display:
        return f"<m:oMathPara {M_NS}>{math}</m:oMathPara>"
    return math


OMML_FRACTION = (
    f"<m:oMath {M_NS}>"
    "<m:f>"
    "<m:num><m:r><m:t>a+b</m:t></m:r></m:num>"
    "<m:den><m:r><m:t>c+d</m:t></m:r></m:den>"
    "</m:f>"
    "</m:oMath>"
)


def add_common(doc: Document) -> None:
    doc.add_heading("Chapter One: Introduction", level=1)
    doc.add_heading("Background", level=2)

    p1 = doc.add_paragraph(
        "Machine learning is a subset of artificial intelligence. "
        "It focuses on algorithms that learn patterns from data."
    )
    assert p1 is not None

    p2 = doc.add_paragraph()
    r1 = p2.add_run("This sentence contains ")
    r1.bold = True
    r2 = p2.add_run("italic and ")
    r2.italic = True
    r3 = p2.add_run("underlined")
    r3.underline = True
    p2.add_run(" inline formatting mixed together.")

    doc.add_paragraph(
        "In quantum mechanics, the wave function describes the quantum state "
        "of a system. Measurement collapses the wave function."
    )

    doc.add_paragraph("First ordered item about apples.", style="List Number")
    doc.add_paragraph("Second ordered item about oranges.", style="List Number")
    doc.add_paragraph("Bullet item about physics.", style="List Bullet")
    doc.add_paragraph("Bullet item about chemistry.", style="List Bullet")

    doc.add_paragraph("")  # 空段落

    centered = doc.add_paragraph("This paragraph is centered for alignment testing.")
    centered.alignment = WD_ALIGN_PARAGRAPH.CENTER


def build_simple() -> None:
    doc = Document()
    add_common(doc)
    doc.save(os.path.join(OUT_DIR, "simple.docx"))


def build_rich() -> None:
    doc = Document()
    add_common(doc)

    # 块级图片（独立段落）
    png1 = make_png(100, 60, (70, 130, 180))
    doc.add_picture(io.BytesIO(png1), width=Inches(1.0))

    doc.add_paragraph("The inline picture below appears inside a sentence.")

    # 行内图片（run 级）
    p_inline = doc.add_paragraph()
    p_inline.add_run("Before the image, then ")
    p_inline.add_run().add_picture(io.BytesIO(make_png(40, 40, (178, 34, 34))))
    p_inline.add_run(" and the sentence continues after it.")

    # 行内公式（分数）
    p_frac = doc.add_paragraph()
    p_frac.add_run("Consider the fraction ")
    p_frac._p.append(parse_xml(OMML_FRACTION))
    p_frac.add_run(" as a simple example of inline math.")

    # 独立段公式（display）
    doc.add_paragraph("The most famous equation follows:")
    p_eq = doc.add_paragraph()
    p_eq._p.append(parse_xml(omml_e_mc2(display=True)))

    # 3x3 表格
    table = doc.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    header = ["Name", "Age", "City"]
    for cell, text in zip(table.rows[0].cells, header):
        cell.text = text
    data = [
        ["Alice", "30", "Beijing"],
        ["Bob", "25", "Shanghai"],
    ]
    for r, row in enumerate(data, start=1):
        for c, text in enumerate(row):
            table.rows[r].cells[c].text = text

    doc.add_paragraph("Final paragraph after the table.")

    doc.save(os.path.join(OUT_DIR, "rich.docx"))


def build_math_para() -> None:
    """v2.3 B4b：含 $...$ 字面文本的段落（MinerU PDF 路径的行内公式形态，
    DOCX 路径通常为 OMML→占位符；此夹具专测 math_slot 防线）"""
    doc = Document()
    doc.add_paragraph("MathPara Fixture")
    doc.add_paragraph("The training loss is $J(\\phi)$ and the rate is $\\eta = 0.1$ per step.")
    doc.save(os.path.join(OUT_DIR, "math_para.docx"))


def main() -> None:
    os.makedirs(OUT_DIR, exist_ok=True)
    build_simple()
    build_rich()
    build_math_para()
    print("fixtures written to", os.path.abspath(OUT_DIR))


if __name__ == "__main__":
    main()
